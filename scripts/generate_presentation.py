#!/usr/bin/env python3
"""
Generates an executive, academic slide presentation (12 widescreen 16:9 slides)
for Molecular Graph Classification Using the Weisfeiler-Lehman Subtree Kernel.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


def build_presentation(output_path: str = "reports/Molecular_Graph_Classification_Presentation.pptx"):
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette: Deep Academic Navy & Clean Accents
    C_NAVY = RGBColor(15, 34, 64)       # Primary dark
    C_BLUE = RGBColor(37, 99, 235)      # Primary accent
    C_DARK = RGBColor(30, 41, 59)       # Body text
    C_MUTED = RGBColor(100, 116, 139)   # Subtitles/secondary
    C_BG_CARD = RGBColor(241, 245, 249) # Card background
    C_WHITE = RGBColor(255, 255, 255)   # White
    C_ACCENT_BG = RGBColor(238, 242, 255)

    def add_header(slide, title_text, category_text="MACHINE LEARNING WITH GRAPHS (23CSE479)"):
        # Category / Kicker
        tb_k = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_k = tb_k.text_frame
        tf_k.word_wrap = True
        p_k = tf_k.paragraphs[0]
        p_k.text = category_text.upper()
        p_k.font.size = Pt(11)
        p_k.font.bold = True
        p_k.font.color.rgb = C_BLUE

        # Main Title
        tb_t = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = C_NAVY

        # Subtle divider
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(226, 232, 240)
        line.line.color.rgb = RGBColor(226, 232, 240)

    # -------------------------------------------------------------------------
    # Slide 1: Title Slide (Dark Navy Background)
    # -------------------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY
    bg1.line.color.rgb = C_NAVY

    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p_kicker = tf1.paragraphs[0]
    p_kicker.text = "MACHINE LEARNING WITH GRAPHS | COURSE CODE: 23CSE479"
    p_kicker.font.size = Pt(13)
    p_kicker.font.bold = True
    p_kicker.font.color.rgb = RGBColor(147, 197, 253)
    p_kicker.space_after = Pt(14)

    p_main = tf1.add_paragraph()
    p_main.text = "Molecular Graph Classification Using the\nWeisfeiler–Lehman Subtree Kernel"
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = C_WHITE
    p_main.space_after = Pt(16)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Comparative Benchmark against Graph Convolutional Networks (GCN) on the NCI1 Dataset"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = RGBColor(203, 213, 225)
    p_sub.space_after = Pt(28)

    p_auth = tf1.add_paragraph()
    p_auth.text = "Author: Vishal Karthikeyan S  |  GitHub: github.com/Vishal-Karthikeyan-S\nAlgorithm Implemented Strictly from Scratch  |  Reproducible Benchmarks & Interactive Dashboard"
    p_auth.font.size = Pt(14)
    p_auth.font.color.rgb = RGBColor(148, 163, 184)

    # -------------------------------------------------------------------------
    # Slide 2: Problem Statement & Motivation
    # -------------------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Problem Formulation: Chemoinformatics & Graph Representation")

    col_w = Inches(3.6)
    cards = [
        ("1. The Molecular Domain",
         "Molecules are naturally represented as undirected attributed graphs:\n\n"
         "• Nodes: Atoms with categorical labels (C, N, O, S, halogens)\n"
         "• Edges: Chemical covalent bonds\n"
         "• Challenge: Invariance to atom ordering (graph isomorphism) and variable molecular sizes."),
        ("2. The NCI1 Benchmark",
         "National Cancer Institute screen for anti-cancer chemical activity:\n\n"
         "• 4,110 chemical compounds\n"
         "• Target: Binary classification (Active cancer inhibitor vs. Inactive)\n"
         "• Highly balanced: 2,053 Class 0 / 2,057 Class 1\n"
         "• Mean size: ~29.9 atoms per molecule."),
        ("3. Core Research Question",
         "Can classical deterministic subtree kernels outperform modern deep Graph Neural Networks?\n\n"
         "• Compare 1-WL Subtree Kernel (from scratch)\n"
         "• Conventional graph-level feature baseline\n"
         "• 3-Layer Graph Convolutional Network (GCN)")
    ]

    for i, (head, body) in enumerate(cards):
        left = Inches(0.8 + i * 4.0)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), col_w, Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = C_BG_CARD
        card.line.color.rgb = RGBColor(203, 213, 225)

        tb = s2.shapes.add_textbox(left + Inches(0.2), Inches(2.0), col_w - Inches(0.4), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = head
        p_h.font.size = Pt(16)
        p_h.font.bold = True
        p_h.font.color.rgb = C_NAVY
        p_h.space_after = Pt(14)

        p_b = tf.add_paragraph()
        p_b.text = body
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = C_DARK

    # -------------------------------------------------------------------------
    # Slide 3: The 1-WL Subtree Kernel Algorithm
    # -------------------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Mathematical Foundations: 1-WL Subtree Kernel from Scratch")

    # Left: Algorithm Steps
    left_card = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(6.8), Inches(4.8))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = C_WHITE
    left_card.line.color.rgb = RGBColor(226, 232, 240)

    tb_alg = s3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(6.4), Inches(4.4))
    tf_alg = tb_alg.text_frame
    tf_alg.word_wrap = True
    p1 = tf_alg.paragraphs[0]
    p1.text = "Step-by-Step 1-WL Algorithm (Shervashidze et al., 2011):"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = C_NAVY
    p1.space_after = Pt(8)

    steps_list = [
        "1. Initialization (t = 0): Assign each atom its categorical label ℓ₀(v) = atom_type(v).",
        "2. Neighborhood Multiset Aggregation: For each atom v, gather neighbor labels:\n   𝓡_t(v) = { ℓ_{t-1}(u) : u ∈ 𝒩(v) }",
        "3. Deterministic Sorting & Hash Compression:\n   ℓ_t(v) = hash( t, ℓ_{t-1}(v), sort(𝓡_t(v)) )",
        "4. Subtree Histogram Counting: Compute occurrence counts of all discovered subtree labels:\n   ϕ_t(G) = [ c(G, σ₁)  c(G, σ₂) … c(G, σ_{|Σ_t|}) ]ᵀ",
        "5. Kernel Gram Matrix: Evaluate inner products with cosine normalization:\n   K_norm(Gᵢ, Gⱼ) = ⟨ϕ(Gᵢ), ϕ(Gⱼ)⟩ / ( ‖ϕ(Gᵢ)‖₂ · ‖ϕ(Gⱼ)‖₂ )"
    ]
    for step in steps_list:
        p_step = tf_alg.add_paragraph()
        p_step.text = step
        p_step.font.size = Pt(11.5)
        p_step.font.color.rgb = C_DARK
        p_step.space_after = Pt(6)

    # Right: Complexity & Implementation Box
    right_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.8), Inches(4.7), Inches(4.8))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = C_ACCENT_BG
    right_card.line.color.rgb = RGBColor(199, 210, 254)

    tb_comp = s3.shapes.add_textbox(Inches(8.0), Inches(2.0), Inches(4.3), Inches(4.4))
    tf_comp = tb_comp.text_frame
    tf_comp.word_wrap = True
    pc1 = tf_comp.paragraphs[0]
    pc1.text = "Implementation & Complexity:"
    pc1.font.size = Pt(16)
    pc1.font.bold = True
    pc1.font.color.rgb = C_BLUE
    pc1.space_after = Pt(10)

    comp_items = [
        ("• Scratch Implementation:", "Strictly manual in src/wl_kernel.py without GraKeL or NetworkX kernel tools."),
        ("• Time Complexity:", "𝒪(h · N · m) for feature extraction across N graphs with m edges."),
        ("• Memory & Sparse Matrix:", "Uses scipy.sparse.csr_matrix with space bounded by 𝒪(N · n · h)."),
        ("• Mercer Guarantee:", "K is proven symmetric (K = Kᵀ) and positive semi-definite (λᵢ ≥ 0)."),
        ("• Zero Data Leakage:", "Vocabulary fitted solely on train graphs; K_test_train (N_test × N_train) evaluated properly.")
    ]
    for h_txt, b_txt in comp_items:
        p_c = tf_comp.add_paragraph()
        p_c.text = f"{h_txt} {b_txt}"
        p_c.font.size = Pt(11.5)
        p_c.font.color.rgb = C_DARK
        p_c.space_after = Pt(5)

    # -------------------------------------------------------------------------
    # Slide 4: Hand-Crafted Proof Example (Cycle C3 vs Path P3)
    # -------------------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Algorithmic Walkthrough: Distinguishing Cycle C3 vs Path P3")

    c1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = C_BG_CARD
    c1.line.color.rgb = RGBColor(203, 213, 225)

    tb_c1 = s4.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    p_t1 = tf_c1.paragraphs[0]
    p_t1.text = "The Setup: Identical Node Counts & Labels"
    p_t1.font.size = Pt(16)
    p_t1.font.bold = True
    p_t1.font.color.rgb = C_NAVY
    p_t1.space_after = Pt(8)

    p_b1 = tf_c1.add_paragraph()
    p_b1.text = (
        "Consider two distinct graphs with 3 identical atoms [1, 1, 1]:\n\n"
        "• Graph A (C₃ Cycle): Triangle edges (0-1, 1-2, 2-0)\n"
        "  Every atom has degree = 2.\n\n"
        "• Graph B (P₃ Path): Linear chain edges (0-1, 1-2)\n"
        "  Endpoint atoms have degree = 1; center atom has degree = 2.\n\n"
        "At Iteration h = 0:\n"
        "Both graphs produce identical feature vectors: ϕ₀ = [3.0]\n"
        "A naive 'bag of atoms' model cannot tell them apart!"
    )
    p_b1.font.size = Pt(12)
    p_b1.font.color.rgb = C_DARK

    c2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = C_WHITE
    c2.line.color.rgb = RGBColor(203, 213, 225)

    tb_c2 = s4.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    p_t2 = tf_c2.paragraphs[0]
    p_t2.text = "At Iteration h = 1: Topological Partitioning"
    p_t2.font.size = Pt(16)
    p_t2.font.bold = True
    p_t2.font.color.rgb = C_BLUE
    p_t2.space_after = Pt(8)

    p_b2 = tf_c2.add_paragraph()
    p_b2.text = (
        "Neighborhood Signatures at h = 1:\n\n"
        "• For C₃: All 3 nodes have signature (1, 0, (0, 0)) → ID 1\n"
        "  ϕ(C₃) = [3, 3, 0]ᵀ\n\n"
        "• For P₃: Center node signature (1, 0, (0, 0)) → ID 1\n"
        "  Two endpoints signature (1, 0, (0,)) → ID 2\n"
        "  ϕ(P₃) = [3, 1, 2]ᵀ\n\n"
        "Resulting Kernel Gram Matrix:\n"
        "  K = ⎡ 18   12 ⎤\n"
        "      ⎣ 12   14 ⎦\n\n"
        "Normalized similarity: 12 / √(18 · 14) = 0.7559 < 1.0\n"
        "1-WL successfully separates the cycle from the path!"
    )
    p_b2.font.size = Pt(12)
    p_b2.font.color.rgb = C_DARK

    # -------------------------------------------------------------------------
    # Slide 5: System Architecture & Workflow Pipeline
    # -------------------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "End-to-End System Pipeline & Architecture")

    blocks = [
        ("1. Ingestion", "NCI1 Dataset\n• 4,110 molecules\n• PyG TUDataset\n• Atom & bond parsing"),
        ("2. Preprocessing", "MolecularGraph\n• Node label mapping\n• Adjacency lists\n• Data validation"),
        ("3. WL Kernel", "Subtree Relabeling\n• Depths h=0,1,2,3\n• Sparse histograms\n• Gram matrix K"),
        ("4. Learning", "Precomputed SVM\n• Strict train-test\n• Decision margins\n• Grid search C"),
        ("5. Baselines", "Fair Comparison\n• Graph features SVM\n• 3-layer GCN (PyG)\n• Early stopping")
    ]

    for i, (head, body) in enumerate(blocks):
        left = Inches(0.8 + i * 2.4)
        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.2), Inches(2.2), Inches(3.8))
        c.fill.solid()
        c.fill.fore_color.rgb = C_BG_CARD if i % 2 == 0 else C_ACCENT_BG
        c.line.color.rgb = RGBColor(203, 213, 225)

        tb = s5.shapes.add_textbox(left + Inches(0.1), Inches(2.4), Inches(2.0), Inches(3.4))
        tf = tb.text_frame
        tf.word_wrap = True
        ph = tf.paragraphs[0]
        ph.text = head
        ph.font.size = Pt(15)
        ph.font.bold = True
        ph.font.color.rgb = C_NAVY
        ph.space_after = Pt(12)

        pb = tf.add_paragraph()
        pb.text = body
        pb.font.size = Pt(12)
        pb.font.color.rgb = C_DARK

    # -------------------------------------------------------------------------
    # Slide 6: Empirical Benchmark Results
    # -------------------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Empirical Benchmark Results on NCI1")

    # Table on Left
    rows = 7
    cols = 6
    top = Inches(1.8)
    left = Inches(0.8)
    width = Inches(11.7)
    height = Inches(4.8)

    table_shape = s6.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    # Column widths
    table.columns[0].width = Inches(3.7)
    table.columns[1].width = Inches(1.6)
    table.columns[2].width = Inches(1.6)
    table.columns[3].width = Inches(1.6)
    table.columns[4].width = Inches(1.6)
    table.columns[5].width = Inches(1.6)

    headers = ["Model Architecture", "Accuracy", "F1-Score", "ROC-AUC", "Train Time", "Features"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_WHITE

    data = [
        ("Graph Feature Baseline + SVM", "66.25%", "0.6494", "0.7125", "0.042 s", "49"),
        ("WL Subtree Kernel (h=0) + SVM", "52.50%", "0.5929", "0.5823", "0.021 s", "26"),
        ("WL Subtree Kernel (h=1) + SVM", "65.83%", "0.6273", "0.7078", "0.042 s", "206"),
        ("WL Subtree Kernel (h=2) + SVM", "71.25%", "0.6849", "0.7685", "0.062 s", "2,162"),
        ("WL Subtree Kernel (h=3) + SVM", "75.00%", "0.7321", "0.7881", "0.083 s", "10,282"),
        ("3-Layer GCN (PyG Baseline)", "67.08%", "0.6638", "0.7594", "41.307 s", "Latent (64-d)")
    ]

    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = val
            cell.fill.solid()
            # Highlight top performer (row index 4)
            if i == 4:
                cell.fill.fore_color.rgb = RGBColor(220, 252, 231)  # Soft green
            elif i % 2 == 1:
                cell.fill.fore_color.rgb = C_BG_CARD
            else:
                cell.fill.fore_color.rgb = C_WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = C_NAVY if i == 4 else C_DARK
            if i == 4 or j == 0:
                p.font.bold = True

    # -------------------------------------------------------------------------
    # Slide 7: WL Subtree Depth Analysis
    # -------------------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Effect of Subtree Depth (h=0 to h=3) on Performance")

    c1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = C_BG_CARD
    c1.line.color.rgb = RGBColor(203, 213, 225)

    tb_d1 = s7.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_d1 = tb_d1.text_frame
    tf_d1.word_wrap = True
    p_th1 = tf_d1.paragraphs[0]
    p_th1.text = "Subtree Horizon & Chemical Representation:"
    p_th1.font.size = Pt(16)
    p_th1.font.bold = True
    p_th1.font.color.rgb = C_NAVY
    p_th1.space_after = Pt(8)

    p_bd1 = tf_d1.add_paragraph()
    p_bd1.text = (
        "• h = 0 (Bag-of-Atoms): 52.50% Accuracy\n"
        "  Only counts raw atom types. Disregards connectivity entirely.\n\n"
        "• h = 1 (1-Hop Bond Environments): 65.83% (+13.33%)\n"
        "  Captures immediate functional groups (e.g. carbonyls C=O, hydroxyls -OH).\n\n"
        "• h = 2 (2-Hop Extended Substructures): 71.25% (+5.42%)\n"
        "  Resolves aromatic rings, conjugate bonds, and branched chains.\n\n"
        "• h = 3 (3-Hop Pharmacophore Motifs): 75.00% (+3.75%)\n"
        "  Encodes multi-atom spatial pharmacophores critical to cancer inhibition."
    )
    p_bd1.font.size = Pt(12.5)
    p_bd1.font.color.rgb = C_DARK

    c2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = C_WHITE
    c2.line.color.rgb = RGBColor(203, 213, 225)

    tb_d2 = s7.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf_d2 = tb_d2.text_frame
    tf_d2.word_wrap = True
    p_th2 = tf_d2.paragraphs[0]
    p_th2.text = "Vocabulary Growth & Diminishing Returns:"
    p_th2.font.size = Pt(16)
    p_th2.font.bold = True
    p_th2.font.color.rgb = C_BLUE
    p_th2.space_after = Pt(8)

    p_bd2 = tf_d2.add_paragraph()
    p_bd2.text = (
        "Vocabulary Expansion Across Depths:\n"
        "  h=0: 26 features\n"
        "  h=1: 206 features\n"
        "  h=2: 2,162 features\n"
        "  h=3: 10,282 features\n\n"
        "Key Research Insight:\n"
        "• As h increases beyond 3, subtree patterns become increasingly unique to single molecules.\n"
        "• Feature sparsity increases rapidly.\n"
        "• h=3 strikes an optimal trade-off between discriminative power and pattern sharing."
    )
    p_bd2.font.size = Pt(12.5)
    p_bd2.font.color.rgb = C_DARK

    # -------------------------------------------------------------------------
    # Slide 8: WL Kernel vs. Graph Neural Networks
    # -------------------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Theoretical & Computational Comparison: WL Kernel vs GCN")

    cards_comp = [
        ("Weisfeiler–Lehman Subtree Kernel",
         "• Paradigm: Discrete Subtree Enumeration + SVM\n"
         "• Accuracy: 75.00% (Superior on NCI1)\n"
         "• Training Speed: 0.083 seconds (~500x faster)\n"
         "• Optimization: Globally optimal convex quadratic program\n"
         "• Determinism: 100% reproducible; zero stochastic variance\n"
         "• Inductive Bias: Exact discrete subgraph frequency match"),
        ("Graph Convolutional Network (GCN)",
         "• Paradigm: Continuous Neural Message Passing (PyG)\n"
         "• Accuracy: 67.08%\n"
         "• Training Speed: 41.307 seconds (60 epochs)\n"
         "• Optimization: Non-convex SGD / Adam backpropagation\n"
         "• Limitations: Over-smoothing risk across layers; mean pooling bounded by 1-WL expressive limit (Xu et al., 2019)")
    ]

    for i, (title_c, text_c) in enumerate(cards_comp):
        left = Inches(0.8 + i * 5.9)
        c = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(5.6), Inches(4.8))
        c.fill.solid()
        c.fill.fore_color.rgb = C_ACCENT_BG if i == 0 else C_BG_CARD
        c.line.color.rgb = RGBColor(199, 210, 254) if i == 0 else RGBColor(203, 213, 225)

        tb = s8.shapes.add_textbox(left + Inches(0.2), Inches(2.0), Inches(5.2), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title_c
        pt.font.size = Pt(16)
        pt.font.bold = True
        pt.font.color.rgb = C_NAVY
        pt.space_after = Pt(12)

        pb = tf.add_paragraph()
        pb.text = text_c
        pb.font.size = Pt(13)
        pb.font.color.rgb = C_DARK

    # -------------------------------------------------------------------------
    # Slide 9: Interactive Streamlit Web Application
    # -------------------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Interactive User Interface & Streamlit Dashboard")

    features = [
        ("📌 Project Overview", "Displays problem formulation, theoretical background, and key benchmark highlights in real time."),
        ("📊 Dataset Explorer", "Inspects NCI1 statistics, class balance, atom counts, and renders interactive molecular graphs."),
        ("🔬 WL Kernel Explorer", "Step-by-step walkthrough of subtree relabeling on the C3 vs P3 example and live NCI1 compounds."),
        ("📈 Empirical Model Comparison", "Interactive comparison tables, ROC-AUC metrics, confusion matrices, and runtime charts."),
        ("🎯 Live Prediction Demo", "Select any molecule and observe real-time inference, ground truth, and signed SVM decision margin.")
    ]

    for i, (h, b) in enumerate(features):
        top_pos = Inches(1.8 + i * 0.95)
        rect = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.7), Inches(0.82))
        rect.fill.solid()
        rect.fill.fore_color.rgb = C_BG_CARD
        rect.line.color.rgb = RGBColor(226, 232, 240)

        tb = s9.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.08), Inches(11.3), Inches(0.66))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{h}:  {b}"
        p.font.size = Pt(13)
        p.font.color.rgb = C_DARK

    # -------------------------------------------------------------------------
    # Slide 10: Engineering Rigor & Codebase Structure
    # -------------------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Software Engineering & Academic Quality Standards")

    c1 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = C_WHITE
    c1.line.color.rgb = RGBColor(203, 213, 225)

    tb_s1 = s10.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_s1 = tb_s1.text_frame
    tf_s1.word_wrap = True
    pt1 = tf_s1.paragraphs[0]
    pt1.text = "Modular Architecture & Separation of Concerns"
    pt1.font.size = Pt(16)
    pt1.font.bold = True
    pt1.font.color.rgb = C_NAVY
    pt1.space_after = Pt(8)

    pb1 = tf_s1.add_paragraph()
    pb1.text = (
        "• src/data_loader.py: Automated download, fallback mirrors, inspection\n"
        "• src/preprocessing.py: Validated MolecularGraph dataclass\n"
        "• src/wl_kernel.py: 1-WL subtree kernel from scratch\n"
        "• src/wl_svm.py: Decoupled train-train & test-train Gram matrices\n"
        "• src/baseline.py: Non-structural feature extraction & SVM\n"
        "• src/gcn_model.py: 3-layer GCN with PyTorch Geometric\n"
        "• src/experiments.py: Automated benchmark harness\n"
        "• src/analysis.py & src/visualization.py: Reporting and plots"
    )
    pb1.font.size = Pt(11.5)
    pb1.font.color.rgb = C_DARK

    c2 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = C_BG_CARD
    c2.line.color.rgb = RGBColor(203, 213, 225)

    tb_s2 = s10.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf_s2 = tb_s2.text_frame
    tf_s2.word_wrap = True
    pt2 = tf_s2.paragraphs[0]
    pt2.text = "Automated Unit Tests & Reproducibility"
    pt2.font.size = Pt(16)
    pt2.font.bold = True
    pt2.font.color.rgb = C_BLUE
    pt2.space_after = Pt(8)

    pb2 = tf_s2.add_paragraph()
    pb2.text = (
        "• 7 Passing Unit Tests (Pytest):\n"
        "  - Label updates & feature dimensions\n"
        "  - Mathematical kernel symmetry K = K^T\n"
        "  - Nonnegative & normalized diagonal entries\n"
        "  - Mercer positive semi-definiteness (eigenvalues >= 0)\n"
        "  - Correct train-train & test-train Gram shapes\n"
        "  - Deterministic feature extraction\n"
        "  - Hand-constructed C3 vs P3 separation"
    )
    pb2.font.size = Pt(12.5)
    pb2.font.color.rgb = C_DARK

    # -------------------------------------------------------------------------
    # Slide 11: Deployment & GitHub Integration
    # -------------------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Deployment Workflow & GitHub Integration")

    c1 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = C_ACCENT_BG
    c1.line.color.rgb = RGBColor(199, 210, 254)

    tb_g1 = s11.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_g1 = tb_g1.text_frame
    tf_g1.word_wrap = True
    pt1 = tf_g1.paragraphs[0]
    pt1.text = "GitHub Repository Setup"
    pt1.font.size = Pt(16)
    pt1.font.bold = True
    pt1.font.color.rgb = C_NAVY
    pt1.space_after = Pt(8)

    pb1 = tf_g1.add_paragraph()
    pb1.text = (
        "User GitHub: github.com/Vishal-Karthikeyan-S\n\n"
        "Steps to publish to GitHub:\n"
        "1. Initialize Git repository:\n"
        "   git init\n"
        "2. Add remote origin:\n"
        "   git remote add origin https://github.com/Vishal-Karthikeyan-S/Molecular-Graph-Classification-WL-Kernel.git\n"
        "3. Commit and push:\n"
        "   git add .\n"
        "   git commit -m 'Initial commit: Molecular graph classification with WL kernel'\n"
        "   git branch -M main\n"
        "   git push -u origin main"
    )
    pb1.font.size = Pt(12)
    pb1.font.color.rgb = C_DARK

    c2 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = C_WHITE
    c2.line.color.rgb = RGBColor(203, 213, 225)

    tb_g2 = s11.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf_g2 = tb_g2.text_frame
    tf_g2.word_wrap = True
    pt2 = tf_g2.paragraphs[0]
    pt2.text = "Streamlit Community Cloud Deployment"
    pt2.font.size = Pt(16)
    pt2.font.bold = True
    pt2.font.color.rgb = C_BLUE
    pt2.space_after = Pt(8)

    pb2 = tf_g2.add_paragraph()
    pb2.text = (
        "Zero-Configuration Deployment on Streamlit Cloud:\n\n"
        "1. Visit share.streamlit.io and sign in with GitHub.\n"
        "2. Click 'New App'.\n"
        "3. Select repository: Vishal-Karthikeyan-S/Molecular-Graph-Classification-WL-Kernel\n"
        "4. Branch: main\n"
        "5. Main file path: app.py\n"
        "6. Click Deploy!\n\n"
        "The application automatically downloads NCI1 raw data and runs with cached performance."
    )
    pb2.font.size = Pt(12)
    pb2.font.color.rgb = C_DARK

    # -------------------------------------------------------------------------
    # Slide 12: Conclusions & Key Takeaways
    # -------------------------------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    bg12 = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg12.fill.solid()
    bg12.fill.fore_color.rgb = C_NAVY
    bg12.line.color.rgb = C_NAVY

    tb12 = s12.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.2))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    p_c1 = tf12.paragraphs[0]
    p_c1.text = "CONCLUSIONS & KEY TAKEAWAYS"
    p_c1.font.size = Pt(14)
    p_c1.font.bold = True
    p_c1.font.color.rgb = RGBColor(147, 197, 253)
    p_c1.space_after = Pt(16)

    p_c2 = tf12.add_paragraph()
    p_c2.text = (
        "1. Superiority of Neighborhood Relabeling:\n"
        "   The 1-WL subtree kernel achieves 75.00% accuracy on NCI1, demonstrating that discrete\n"
        "   subtree patterns represent localized chemical motifs far more effectively than global graph features (66.25%).\n\n"
        "2. Algorithmic Efficiency:\n"
        "   The WL kernel is ~500x faster to train than a standard GCN while yielding higher accuracy,\n"
        "   proving that classical graph kernels remain exceptionally strong baselines in chemoinformatics.\n\n"
        "3. Complete Academic Deliverable:\n"
        "   • From-scratch algorithmic implementation with unit testing\n"
        "   • Rigorous fair benchmarking preventing test data leakage\n"
        "   • Cloud-ready interactive Streamlit interface & complete academic documentation"
    )
    p_c2.font.size = Pt(16)
    p_c2.font.color.rgb = C_WHITE
    p_c2.space_after = Pt(24)

    p_c3 = tf12.add_paragraph()
    p_c3.text = "Thank You!  |  GitHub: github.com/Vishal-Karthikeyan-S"
    p_c3.font.size = Pt(15)
    p_c3.font.bold = True
    p_c3.font.color.rgb = RGBColor(203, 213, 225)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at {output_path}")


if __name__ == "__main__":
    build_presentation()
